import openpyxl
import os

print('''
      [1] Criar uma nova planilha
      [2] Editar ou buscar em uma planilha existente
     ''')

tabela = input('Escolha uma opção acima: ')

match tabela:
    case '1':  ################ CRIAÇÃO DA PLANILHA EXCEL ##################
        print(' ')
        nome_arquivo = input("Qual o nome da nova planilha (ex: tabela.xlsx): ")
        
        # Garante que o arquivo terá a extensão .xlsx
        if not nome_arquivo.endswith('.xlsx'):
            nome_arquivo += '.xlsx'

        # Cria a pasta de trabalho e seleciona a aba ativa
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dados"

        # Adiciona os cabeçalhos das colunas
        campos = ["Nome", "Idade", "Cidade"]
        ws.append(campos)

        n1 = int(input('Qual a quantidade de registros iniciais: '))
        contador = 1

        while contador <= n1:
            print(f"\n--- Inserindo registro {contador} de {n1} ---")
            nome = input('Nome: ')
            idd = int(input('Idade: '))
            cdd = input('Cidade: ')

            # Adiciona uma nova linha com os dados
            ws.append([nome, idd, cdd])
            contador += 1

        # Salva o arquivo no formato Excel nativo
        wb.save(nome_arquivo)
        print(f'\nArquivo "{nome_arquivo}" criado com sucesso!')

    case '2':  ################ ADICIONAR / PROCURAR ###############################
        print('''
      [1] Adicionar um novo registro
      [2] Procurar um nome na planilha   
               ''')
        opcao = input('Escolha uma opção acima: ')
        print('')

        match opcao:
            case '1':  ################ ADICIONAR NOVO REGISTRO ##################
                nome_arquivo = input("Qual o nome da planilha que deseja editar (ex: tabela.xlsx): ")
                if not nome_arquivo.endswith('.xlsx'):
                    nome_arquivo += '.xlsx'

                try:
                    # Carrega a planilha existente
                    wb = openpyxl.load_workbook(nome_arquivo)
                    ws = wb.active

                    nome = input('Qual nome? ')
                    idd = int(input('Qual idade? '))
                    cdd = input('Qual cidade nasceu? ')

                    # Adiciona a linha no final da tabela
                    ws.append([nome, idd, cdd])
                    wb.save(nome_arquivo)
                    print("\nRegistro adicionado com sucesso no arquivo Excel!")

                except FileNotFoundError:
                    print("\nArquivo não encontrado. Crie a tabela primeiro usando a opção [1].")

            case '2':  ################ PROCURAR REGISTRO ##################
                nome_arquivo = input("Qual o nome da planilha que deseja pesquisar (ex: tabela.xlsx): ")
                if not nome_arquivo.endswith('.xlsx'):
                    nome_arquivo += '.xlsx'

                nome_procurado = input('Qual nome você quer procurar? ')

                try:
                    # Carrega o arquivo Excel
                    wb = openpyxl.load_workbook(nome_arquivo)
                    ws = wb.active

                    # Lê a primeira linha para obter os cabeçalhos das colunas
                    header = [cell.value for cell in ws[1]]
                    encontrado = False

                    # Percorre as linhas do Excel a partir da linha 2 (ignorando o cabeçalho)
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not row or row[0] is None:
                            continue

                        # O primeiro elemento (índice 0) é a coluna "Nome"
                        nome_na_celula = str(row[0]).strip().lower()
                        if nome_na_celula == nome_procurado.strip().lower():
                            print("\nRegistro encontrado:")
                            for chave, valor in zip(header, row):
                                print(f"  {chave}: {valor}")
                            encontrado = True

                    if not encontrado:
                        print("\nNome não encontrado na planilha.")

                except FileNotFoundError:
                    print("\nArquivo não encontrado. Verifique o nome da tabela.")

            case _:
                print("Opção inválida.")

    case _:
        print("Opção inválida.")